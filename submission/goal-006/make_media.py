"""Offline Goal 006 editor: historical source is read-only, all outputs stay in Goal 006."""
from pathlib import Path
import array, hashlib, json, math, re, shutil, subprocess, textwrap, wave, zipfile
from functools import lru_cache
from PIL import Image, ImageDraw, ImageFont, ImageOps
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'submission/goal-006'
WORK = ROOT / '.validation/goal-006-media'
ORIGINAL = ROOT.parent / 'talk-me-home'
SOURCE = ORIGINAL / '.validation/goal-005-media/Goal005_Attempt12_Success_Uncut.mp4'
RAW = ORIGINAL / '.validation/goal-004c-live/2026-09-29T11-10-24-116Z-voice-mission'
EXPECTED = 'd69a0b9a1e793294399e8731cffb7a4c8765772e685510598ff8b7d04e8d4812'
BUILD = '2edf7914a008143843923b04a9bf3a1fe41f1f68'
FONT = '/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf'
BOLD = '/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf'
SERIF = '/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf'
PAPER, INK, GREEN, GOLD, MUTED = 'F4EFE2','24382F','557862','B58A46','607066'
for p in (OUT, WORK, WORK/'cards', WORK/'edit', WORK/'review', WORK/'slides', OUT/'screenshots'): p.mkdir(parents=True, exist_ok=True)
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def run(a): subprocess.run(a, check=True, stdout=subprocess.DEVNULL)
def probe(p): return json.loads(subprocess.check_output(['ffprobe','-v','error','-show_format','-show_streams','-of','json',str(p)]))
def fnt(size,bold=False,serif=False): return ImageFont.truetype(SERIF if serif else BOLD if bold else FONT,size)
def savej(p,j): p.write_text(json.dumps(j,indent=2)+'\n')
assert sha(SOURCE)==EXPECTED, 'Historical source hash changed; do not silently refresh.'
PRACTICE_PATH=OUT/'practice-source.json'
PRACTICE=json.loads(PRACTICE_PATH.read_text()) if PRACTICE_PATH.exists() else None
if PRACTICE:
 assert PRACTICE['mode']=='Practice — deterministic simulation — new gameplay build'
 assert PRACTICE['realProviderCalls']==0 and PRACTICE['recorderSecuredAtHome'] is True
 for item in [PRACTICE['video'],PRACTICE['captureReceipt'],PRACTICE['screenshotCaptureReceipt'],*PRACTICE['screenshots']]:
  path=(ROOT/item['path']).resolve()
  assert path.is_relative_to(ROOT) and path.is_file() and sha(path)==item['sha256'], 'Review changed Practice source before rendering.'

def card(name, eyebrow, title, body, footer=''):
 im=Image.new('RGB',(1920,1080),'#'+INK); d=ImageDraw.Draw(im)
 d.rectangle((110,105,190,111),fill='#'+GOLD)
 d.text((110,150),eyebrow.upper(),font=fnt(25,True),fill='#C9D3BD')
 d.multiline_text((110,280),title,font=fnt(78,serif=True),spacing=14,fill='#'+PAPER)
 d.multiline_text((115,560),body,font=fnt(38),spacing=22,fill='#E5DECA')
 if footer:d.text((115,965),footer,font=fnt(24),fill='#C9D3BD')
 im.save(WORK/'cards'/f'{name}.png')

card('premise','Talk Me Home','You have the map.\nPip has eyes and hands.','Read the document. Share what you know.\nDecide together. Confirm the exact action.','Recorded gameplay follows: real AssemblyAI, synthetic microphone input.')
card('technology','Voice makes the partner possible','Speech carries the cooperation.','Microphone → AssemblyAI → agent tools → spoken reply\nThe server validates. Your console confirms.','Conversation connects the partners. Physical actions still need your exact confirmation.')
card('status','A working prototype','One rescue demonstrated.\nReliability still in development.','A second run on the same build stalled.\nLocal prototype; public access and submission are not yet available.','Edited from one original run. Synthetic input + digital playback, not a human playtest.')
if PRACTICE:
 card('status','A working prototype','One historical real rescue.\nNew choices, tested in Practice.','Repeat-run Voice reliability remains in development.\nLocal prototype; hosting and submission remain separate.','Historical Voice uses synthetic input. New Practice is simulation, not a new Live pass.')

segments=[
 dict(kind='voice',start=3.7,duration=6.5,chapter='A PARTNER NEEDS YOUR HELP',cut='Original greeting'),
 dict(kind='card',card='premise',duration=7),
 dict(kind='voice',start=38.1,duration=14,chapter='01  CARGO / A SHARED SUPPLY',cut='Cut to Cargo: read the document, agree on the Latch'),
 dict(kind='voice',start=62.3,duration=40.4,chapter='01  CARGO / ONE SHARED SUCCESS',cut='Cut forward: Latch confirmed → Power off → crossing'),
 dict(kind='voice',start=222.2,duration=35.1,chapter='02  GALLERY / A ROUTE IS A HYPOTHESIS',cut='Later: an open gate is blocked; check the alternative'),
 dict(kind='voice',start=259.5,duration=27,chapter='02  GALLERY / CHANGE THE PLAN',cut='Same conversation: choose southwest and confirm the backtrack'),
 dict(kind='voice',start=428.3,duration=19.3,chapter='03  DOCK / TWO PARTNERS, ONE TRANSFER',cut='Later at Dock: Pip holds the contact after exact confirmation'),
 dict(kind='voice',start=462.1,duration=4.2,chapter='03  DOCK / STORE THE ENERGY',cut='Cut forward: Mission Control charges and stores energy'),
 dict(kind='voice',start=507.8,duration=5.4,chapter='03  DOCK / READY TO RETURN',cut='Later: contact released and boarding confirmed'),
 dict(kind='voice',start=555.25,duration=29,chapter='HOME / THE CONFIRMED RETURN',cut='Later: owner grants return; one final exact confirmation'),
 dict(kind='card',card='technology',duration=8),
 dict(kind='card',card='status',duration=7),
]
if PRACTICE:
 # Remove only two silent Cargo waits. Preserve every selected utterance and the
 # uninterrupted 35.1-second Gallery question/report/alternative-check exchange.
 segments[3:4]=[
  dict(kind='voice',start=62.3,duration=7.7,chapter='01  CARGO / ONE SHARED SUCCESS',cut='Cut forward: Latch confirmed → Mission Control turns Power off'),
  dict(kind='voice',start=72.7,duration=11.3,chapter='01  CARGO / ONE SHARED SUCCESS',cut='Waiting shortened: the Conveyor stops; choose the crossing'),
  dict(kind='voice',start=86.7,duration=16,chapter='01  CARGO / ONE SHARED SUCCESS',cut='Waiting shortened: exact crossing proposal, confirmation and arrival'),
 ]
 cuts=PRACTICE['cuts'];assert abs(sum(c['duration'] for c in cuts)-12)<.000001
 insert=[dict(kind='practice',start=c['start'],duration=c['duration'],chapter=c['chapter'],cut=c['caption'],source=PRACTICE['video']['path']) for c in cuts]
 segments[-2:-2]=insert
 assert abs(sum(s['duration'] for s in segments)-209.5)<.000001

@lru_cache(maxsize=1)
def transcripts():
 events=json.loads((RAW/'audio-evidence.json').read_text())['events']; last=0; speech=[]
 for e in events:
  typ=e['type']; end=e['atMs']/1000
  if typ=='transcript.agent':
   onset=[x['atMs']/1000 for x in events if x['type']=='audio.rendered.onset' and last+.1<x['atMs']/1000<=end]
   if onset: speech.append(dict(speaker='Pip',start=min(onset)+.0802,end=end+.36,text=e['text']))
   last=end
  elif typ=='transcript.user':
   starts=[x['atMs']/1000 for x in events if x['type']=='synthetic.speech.started' and x['atMs']/1000<=end]
   speech.append(dict(speaker='Player (synthetic)',start=starts[-1]+.0802,end=end+.12,text=e['text']))
 # Refine utterance boundaries against the original independent digital channels.
 # This preserves real pauses without equating a transcript-final event with sound ending.
 channels={}
 for speaker,filename in [('Pip','rendered-digital.wav'),('Player (synthetic)','input-digital.wav')]:
  with wave.open(str(RAW/filename)) as w:
   assert w.getnchannels()==1 and w.getsampwidth()==2
   channels[speaker]=(w.getframerate(),array.array('h',w.readframes(w.getnframes())))
 for u in speech:
  rate,samples=channels[u['speaker']]
  lo=max(0,round((u['start']-.0802-.08)*rate));hi=min(len(samples),round((u['end']-.0802+.6)*rate))
  first=next((i for i in range(lo,hi) if abs(samples[i])>64),None)
  final=next((i for i in range(hi-1,lo-1,-1) if abs(samples[i])>64),None)
  if first is not None:
   u['start']=first/rate+.0802;u['end']=(final+1)/rate+.0802+.03
 return speech

def srt_time(t):
 ms=round(t*1000); h,ms=divmod(ms,3600000); m,ms=divmod(ms,60000); s,ms=divmod(ms,1000)
 return f'{h:02}:{m:02}:{s:02},{ms:03}'

def balanced_phrases(value):
 """Even word-boundary chunks; never strand a short last word in its own cue."""
 words=value.split(); count=max(1,math.ceil(len(value)/92)); chunks=[]
 while count>1:
  target=len(' '.join(words))/count
  cut=min(range(1,len(words)-count+2),key=lambda n:abs(len(' '.join(words[:n]))-target))
  chunks.append(' '.join(words[:cut]));words=words[cut:];count-=1
 chunks.append(' '.join(words))
 assert ' '.join(chunks)==value
 return chunks

def subtitle_cues(seg,cursor):
 if seg['kind']!='voice':return []
 out=[]; ss=seg['start']; se=ss+seg['duration']
 for u in transcripts():
  if u['start']>=ss and u['end']<=se:
   chunks=balanced_phrases(u['text'])
   total=sum(len(c) for c in chunks); at=u['start']
   for c in chunks:
    end=at+(u['end']-u['start'])*len(c)/total
    assert len(chunks)==1 or end-at>=1.25, 'Do not split an utterance into unreadably short fragments.'
    wrapped='\n'.join(textwrap.wrap(c,width=53,break_long_words=False,break_on_hyphens=False))
    out.append(dict(start=cursor+at-ss,end=cursor+end-ss,text=f"{u['speaker']}: {wrapped}")); at=end
 return out

def render_video():
 cursor=0; cues=[]; timeline=[]
 for i,s in enumerate(segments):
  end=cursor+s['duration']; target=WORK/'edit'/f'segment-{i:02}.mp4'
  cues+=subtitle_cues(s,cursor)
  args=['ffmpeg','-nostdin','-y','-hide_banner','-loglevel','error']
  if s['kind']=='card':
   card_path=WORK/'cards'/f"{s['card']}.png";render_inputs=[card_path]
   args+=['-loop','1','-i',str(card_path)]
   vf='setsar=1'
  else:
   source=ROOT/s['source'] if s['kind']=='practice' else SOURCE
   args+=['-ss',str(s['start']),'-i',str(source)]
   header=WORK/'edit'/f'{i}-header.txt';header.write_text(s['chapter'])
   label=WORK/'edit'/f'{i}-mode.txt';label.write_text('Practice · deterministic simulation · new gameplay build' if s['kind']=='practice' else 'Historical real Voice · edited | synthetic microphone + UI confirmation')
   cut=WORK/'edit'/f'{i}-cut.txt';cut.write_text(s['cut'])
   render_inputs=[source,header,label,cut]
   composition='scale=1528:930:force_original_aspect_ratio=decrease,pad=1920:1080:(ow-iw)/2:64' if s['kind']=='practice' else 'crop=1440:877:0:0,scale=1528:930,pad=1920:1080:196:64'
   vf=f"{composition}:color=0x{INK},setsar=1,drawtext=fontfile={BOLD}:textfile={header}:fontcolor=0x{PAPER}:fontsize=26:x=54:y=19,drawtext=fontfile={FONT}:textfile={label}:fontcolor=0xDDD8C7:fontsize=22:x=w-tw-54:y=21,drawtext=fontfile={FONT}:textfile={cut}:fontcolor=0xE0D3AF:fontsize=22:x=(w-tw)/2:y=1003:enable='lt(t,3)'"
  args+=['-t',str(s['duration']),'-vf',vf,'-map','0:v:0','-an','-r','30','-fps_mode','cfr','-c:v','libx264','-preset','fast','-crf','21','-pix_fmt','yuv420p','-threads','2',str(target)]
  signature=hashlib.sha256(json.dumps({'args':args,'inputs':[(str(p),sha(p)) for p in render_inputs]}).encode()).hexdigest(); stamp=target.with_suffix('.sha')
  if not target.exists() or not stamp.exists() or stamp.read_text()!=signature:
   run(args);stamp.write_text(signature)
  encoded=probe(target)['streams'][0]; frames=int(encoded['nb_frames'])
  assert abs(frames/30-s['duration'])<.000001, 'Use frame-exact segment durations.'
  timeline.append(dict(**s,outputStart=cursor,outputEnd=end,videoFrames=frames,outputStartSample=round(cursor*48000),audioSamples=round(s['duration']*48000)));cursor=end
 playlist=WORK/'edit/concat.txt';playlist.write_text(''.join(f"file '{WORK/'edit'/f'segment-{i:02}.mp4'}'\n" for i in range(len(segments))))
 concatenated=WORK/'edit/assembled-video.mp4';run(['ffmpeg','-nostdin','-y','-v','error','-f','concat','-safe','0','-i',str(playlist),'-c:v','copy','-an',str(concatenated)])
 # Decode the original AAC once. Assemble an exact PCM timeline, then encode AAC once.
 # Concatenating separately encoded AAC carries a priming packet at every edit boundary.
 decoded=WORK/'edit/source-audio-48000.wav'
 if not decoded.exists():run(['ffmpeg','-nostdin','-y','-v','error','-i',str(SOURCE),'-vn','-ac','1','-ar','48000','-c:a','pcm_s16le',str(decoded)])
 with wave.open(str(decoded)) as w:
  assert w.getframerate()==48000 and w.getnchannels()==1 and w.getsampwidth()==2
  raw=array.array('h',w.readframes(w.getnframes()))
 gain=min(1.8,.95*32767/max(abs(v) for v in raw)); pcm=WORK/'edit/final-timeline.wav'
 with wave.open(str(pcm),'wb') as w:
  w.setnchannels(1);w.setsampwidth(2);w.setframerate(48000)
  for s in timeline:
   count=s['audioSamples']
   if s['kind']=='voice':
    start=round(s['start']*48000); chosen=raw[start:start+count]
    scaled=array.array('h',(round(v*gain) for v in chosen));scaled.extend([0]*(count-len(scaled)));w.writeframes(scaled.tobytes())
   else:w.writeframes(b'\0\0'*count)
 with wave.open(str(pcm)) as w:assert w.getnframes()==round(cursor*48000)
 srt='\n\n'.join(f"{i+1}\n{srt_time(c['start'])} --> {srt_time(c['end'])}\n{c['text']}" for i,c in enumerate(cues))+'\n'
 (OUT/'video-subtitles.srt').write_text(srt)
 # Mux the timed SRT as selectable subtitles. The full game captions also remain in-frame.
 run(['ffmpeg','-nostdin','-y','-v','error','-i',str(concatenated),'-i',str(pcm),'-i',str(OUT/'video-subtitles.srt'),'-map','0:v','-map','1:a','-map','2:0','-c:v','copy','-c:a','aac','-b:a','128k','-ar','48000','-ac','2','-c:s','mov_text','-metadata:s:s:0','language=eng','-t',str(cursor),'-movflags','+faststart',str(OUT/'Talk_Me_Home_Submission.mp4')])
 savej(OUT/'edit-plan.json',dict(source=str(SOURCE),sha256=EXPECTED,build=BUILD,practice=PRACTICE,segments=timeline,durationSeconds=cursor,actualGameplaySeconds=sum(s['duration'] for s in segments if s['kind'] in ('voice','practice')),historicalVoiceSeconds=sum(s['duration'] for s in segments if s['kind']=='voice'),practiceSeconds=sum(s['duration'] for s in segments if s['kind']=='practice'),narration=False,speedChanges=False,audioAssembly=dict(method='Sample-exact 48 kHz PCM, source decoded once, one final AAC encode',pcmSamples=round(cursor*48000),fixedGain=gain,perSegmentAacConcatenation=False,practice='Silent browser capture; no historical Voice audio underneath'),subtitleTiming='Original digital channel speech boundaries, using captured onset windows and retained 80.2 ms source alignment. Balanced phrase chunks, not word-level forced alignment.',subtitles=cues))

def frame(t,path):run(['ffmpeg','-nostdin','-y','-v','error','-ss',str(t),'-i',str(SOURCE),'-frames:v','1',str(path)])
for name,t in [('historical-cargo',65),('historical-gallery',271),('historical-dock',465.8),('historical-home',580)]:
 p=WORK/'review'/f'{name}.png'
 if not p.exists():frame(t,p)

def deck():
 prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
 prs.core_properties.title='Talk Me Home — cooperation through conversation';prs.core_properties.author='Talk Me Home'
 def box(s,x,y,w,h,fill):
  z=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));z.fill.solid();z.fill.fore_color.rgb=RGBColor.from_string(fill);z.line.fill.background();return z
 def text(s,t,x,y,w,h,size=22,color=INK,bold=False,serif=False):
  z=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));f=z.text_frame;f.word_wrap=True;f.margin_left=f.margin_right=f.margin_top=f.margin_bottom=0
  for n,line in enumerate(t.split('\n')):
   p=f.paragraphs[0] if n==0 else f.add_paragraph();p.text=line;p.font.name='Liberation Serif' if serif else 'Liberation Sans';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(8)
  return z
 def pic(s,p,x,y,w,h):
  iw,ih=Image.open(p).size;sc=min(w/iw,h/ih);dw,dh=iw*sc,ih*sc
  return s.shapes.add_picture(str(p),Inches(x+(w-dw)/2),Inches(y+(h-dh)/2),Inches(dw),Inches(dh))
 def slide(n,k,t):
  s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(PAPER)
  box(s,.65,.5,.45,.045,GOLD);text(s,k.upper(),1.24,.4,11,.35,11,MUTED,True);text(s,t,.65,1.05,12,.9,35,serif=True)
  footer='Historical real Voice + new Practice · different builds, separately labelled' if PRACTICE else 'Historical real Voice · synthetic microphone input · local prototype'
  text(s,'Talk Me Home · '+footer,.68,7.08,11,.22,9,MUTED);text(s,f'{n:02} / 06',11.9,7.05,.8,.3,10,MUTED);return s
 art=ROOT/'submission/assets/cover-illustration-1600x900.png'
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(INK)
 cropped=WORK/'cards/cover-art-detail.png';Image.open(art).crop((745,60,1600,870)).save(cropped)
 pic(s,cropped,6.1,0,7.23,7.5);box(s,0,0,6.2,7.5,INK)
 text(s,'COOPERATION THROUGH CONVERSATION',.7,.7,5.4,.4,12,'CDD4C0',True);text(s,'Talk Me\nHome.',.67,1.7,5.5,2,58,PAPER,serif=True)
 text(s,'You have the map.\nPip has eyes and hands.',.73,4.37,5.6,1.1,25,PAPER);text(s,'Neither can escape alone.',.73,6.27,5.4,.55,20,'D7D1B9')
 s.notes_slide.notes_text_frame.text='Original project illustration, not a fake gameplay screenshot. Local prototype. No live URL is supplied.'
 s=slide(2,'The asymmetry','One world, two incomplete views.')
 for x,t,b in [(.7,'YOU / MISSION CONTROL','Private maps and equipment manuals\nRemote Power, Relay and Charge\nA plan — and the final confirmation'),(6.85,'PIP / INSIDE THE STATION','Local eyes and hands\nObserved clues and passage checks\nA specific physical proposal')]:
  box(s,x,2.42,5.78,3.22,'E8E1CF' if x<1 else 'E3E8DB');text(s,t,x+.25,2.7,5.28,.5,15,GREEN,True);text(s,b,x+.25,3.43,5.28,1.8,22)
 text(s,'The shared conversation is the bridge. Neither side gets an automatic complete view.',.77,6.12,11.8,.65,22,GREEN)
 s.notes_slide.notes_text_frame.text='Native editable role diagram. Human private map and notes are not supplied to Pip; hidden local state is not streamed into the human map.'
 s=slide(3,'The decision loop','Listen. Decide. Confirm.')
 pic(s,WORK/'review/historical-gallery.png',.65,2.2,8.7,4.65)
 text(s,'1  LISTEN',9.65,2.48,3,.4,14,GREEN,True);text(s,'“there is cargo\nblocking the opening.”',9.65,2.96,3,1,21)
 text(s,'2  REVISE',9.65,4.13,3,.4,14,GREEN,True);text(s,'Check the southwest\npassage. Backtrack.',9.65,4.56,3,.85,21)
 text(s,'3  CONFIRM',9.65,5.79,3,.4,14,GREEN,True);text(s,'Commit the exact move.',9.65,6.22,3,.4,17)
 s.notes_slide.notes_text_frame.text='Actual attempt12 Gallery frame. Quotation is an excerpt from its original provider transcript. Synthetic microphone input, real AssemblyAI, exact UI confirmation. The recording shows a blocked southeast passage and confirmed southwest backtrack.'
 s=slide(4,'An arc with consequences','Three chapters, one rescue.')
 if PRACTICE:
  for y,n,t,b in [(2.28,'01','Cargo Bay','Discover the shared supply.'),(3.38,'02','Relay Gallery','A reported obstruction changes your route.'),(4.48,'03','Return Dock','Hold, charge, store — then return together.')]:
   text(s,n,.72,y,.55,.4,15,GREEN,True);text(s,t,1.4,y,4.6,.48,24,serif=True);text(s,b,1.4,y+.5,4.65,.55,16)
  home=next(p for p in PRACTICE['screenshots'] if p['name']=='home')
  text(s,'PRACTICE · NEW GAMEPLAY BUILD',6.55,2.25,5.9,.35,13,GREEN,True);pic(s,ROOT/home['path'],6.4,2.8,6.28,3.54)
  text(s,'Optional: bring back Pip’s flight recorder.',.73,6.02,5.8,.49,20,GREEN,True);text(s,'Go home directly, or visit the archive. Only a confirmed pickup changes the shelf.',.73,6.56,11.8,.34,16)
  s.notes_slide.notes_text_frame.text='Actual new-build deterministic Practice screenshot, not real Voice. Source '+home['path']+'; screenshot source commit '+home['sourceCommit']+'. Recorder modifier is explicitly selected, optional, and the ending item reflects an admitted pickup. Historical Voice did not contain this mechanic. Its archive location is authored independently of the blocked gate; some routes already pass through it. Video and later screenshot identities are retained separately in practice-source.json.'
 else:
  for x,n,t,b,img in [(.65,'01','Cargo Bay','Discover the shared supply.\nMake the first crossing.','historical-cargo'),(4.85,'02','Relay Gallery','Test a route against reports.\nRevise it when cargo blocks it.','historical-gallery'),(9.04,'03','Return Dock','Hold, charge and store.\nCoordinate the journey home.','historical-dock')]:
   text(s,n,x,2.2,1,.45,15,GREEN,True);text(s,t,x,2.76,3.6,.55,25,serif=True);pic(s,WORK/'review'/f'{img}.png',x,3.45,3.65,2.1);text(s,b,x,5.85,3.7,.9,17)
  s.notes_slide.notes_text_frame.text='All three screenshots are actual frames from one successful historical Voice run on build '+BUILD+'. The optional flight recorder is not in this baseline recording.'
 s=slide(5,'AssemblyAI contribution','Voice makes the partner possible.')
 labels=[('MICROPHONE','Player speaks'),('ASSEMBLYAI','ASR + agent'),('LOCAL TOOLS','Look / propose'),('GAME SERVER','Validate state'),('SPEECH','Pip replies')]
 for i,(a,b) in enumerate(labels):
  x=.7+i*2.51;box(s,x,2.7,2.14,1.3,'E4E9DC');text(s,a,x+.15,2.92,1.86,.35,12,GREEN,True);text(s,b,x+.15,3.42,1.9,.35,15)
  if i<4:text(s,'→',x+2.2,3.08,.3,.4,22,GOLD)
 box(s,3.15,4.58,7.02,1.26,INK);text(s,'HUMAN CONSOLE / EXACT CONFIRMATION',3.4,4.82,6.5,.4,16,PAPER,True);text(s,'Only an admitted confirmation commits a physical proposal.',3.4,5.29,6.5,.32,15,PAPER)
 text(s,'Conversation exchanges information. Authority remains in the game server.',.8,6.35,11.8,.5,21,GREEN)
 s.notes_slide.notes_text_frame.text='Native editable architecture diagram. AssemblyAI delivers recognition, agent generation, client function-call protocol and speech; browser tools use the authoritative game server. Read-only inspections are immediate. Physical actions require a separate exact owner-bound console decision.'
 s=slide(6,'Evidence and next step','A working prototype, with a clear next step.')
 pic(s,WORK/'review/historical-home.png',.67,2.24,7.4,4.38)
 text(s,'HISTORICAL REAL VOICE',8.43,2.42,4.1,.35,13,GREEN,True);text(s,'One full rescue, synthetic input.\nEnding ACK received.',8.43,2.96,4.1,1.0,20)
 text(s,'NEW BUILD / PRACTICE' if PRACTICE else 'STILL TO PROVE',8.43,4.22,4.1,.35,13,GREEN,True);text(s,'Optional recorder demonstrated.\nNo new Live acceptance.' if PRACTICE else 'Repeat-run reliability.\nHuman enjoyment and physical audio.',8.43,4.71,4.1,.9,20)
 text(s,'NEXT / STILL UNMEASURED',8.43,5.98,4.1,.35,13,GREEN,True);text(s,'Human enjoyment, physical audio, reliability.\nHosting and submission remain separate.',8.43,6.38,4.1,.57,14)
 s.notes_slide.notes_text_frame.text='Attempt12 completed Cargo/Gallery/Dock/home with11 exact confirmations and remote ending ACK. Attempt13 on the same candidate stalled in Gallery without ACK. Two-pass target not met; RELEASE_NOT_LIVE_VERIFIED. No users, revenue or enjoyment scores invented. Future playtest is optional and not a prerequisite for local delivery.'
 prs.save(OUT/'Talk_Me_Home_Pitch.pptx')

def cover():
 art=Image.open(ROOT/'submission/assets/cover-illustration-1600x900.png').convert('RGB').crop((745,0,1600,835))
 im=Image.new('RGB',(1920,1080),'#'+INK);right=ImageOps.fit(art,(1020,1080),centering=(.8,.5));im.paste(right,(900,0));d=ImageDraw.Draw(im);d.rectangle((0,0,935,1080),fill='#'+INK)
 d.text((90,140),'COOPERATION THROUGH CONVERSATION',font=fnt(27,True),fill='#D1D7C3');d.multiline_text((85,315),'Talk Me\nHome.',font=fnt(136,serif=True),spacing=5,fill='#'+PAPER)
 d.multiline_text((95,755),'You have the map.\nPip has eyes and hands.',font=fnt(41),spacing=18,fill='#E5DCC6');d.text((95,1010),'Original product illustration · local prototype',font=fnt(22),fill='#C6CEBC');im.save(OUT/'cover-1920x1080.png')

if __name__=='__main__':
 import sys
 action=sys.argv[1] if len(sys.argv)>1 else 'all'
 if action in ('all','deck'):deck();cover()
 if action in ('all','video'):render_video()
 print('Rendered requested Goal 006 assets:',action)

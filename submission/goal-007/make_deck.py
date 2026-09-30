"""Create six editable local Goal 007 slides and copy the unchanged cover.

Run with the existing submission-tools Python environment. Rendering the saved
PPTX with LibreOffice is a separate documented step; historical recipes are read-only.
"""
from pathlib import Path
import hashlib, json, shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'submission/goal-007'
PAPER='F4EFE2'; INK='24382F'; GREEN='58745D'; GOLD='B58A46'; MUTED='687365'; PALE='E2E5CF'; LIGHT='DCE3D1'
SOURCE='66f1bd206057720275fc9d56a6c9fbe72e409bcd'
COVER_HASH='38aa0826a78448c03c99bf717d44817b5aa57b3133071ff05ff6bf0df740ec91'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
original=ROOT/'submission/goal-006/cover-1920x1080.png'
assert sha(original)==COVER_HASH
shutil.copyfile(original,OUT/'cover-1920x1080.png')
assert sha(OUT/'cover-1920x1080.png')==COVER_HASH
capture=json.loads((OUT/'screenshots/capture-provenance.json').read_text())
assert capture['uiSourceCommit']==SOURCE and capture['providerTokenRequests']==0
assert capture['confirmedHome'] and capture['recorderSecuredAtHome']
for item in capture['screenshots']: assert sha(OUT/item['path'])==item['sha256']
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
prs.core_properties.title='Talk Me Home - Goal 007 local submission components'
prs.core_properties.subject='Six editable slides; hosted access and current Live acceptance pending'
prs.core_properties.author='Talk Me Home'
prs.core_properties.keywords='Goal007, Practice, work in progress, no new Live claim'

def text(s,x,y,w,h,value,size=20,color=INK,bold=False,serif=False):
    box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=box.text_frame
    tf.clear();tf.word_wrap=True;tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(value.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line
        p.font.name='Liberation Serif' if serif else 'Liberation Sans';p.font.size=Pt(size)
        p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(8)
    return box

def rect(s,x,y,w,h,fill,line=None,radius=False):
    box=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    box.fill.solid();box.fill.fore_color.rgb=RGBColor.from_string(fill)
    box.line.fill.background() if not line else None
    if line:box.line.color.rgb=RGBColor.from_string(line)
    return box

def picture(s,path,x,y,w,h,region=None):
    # Native PowerPoint crop preserves original pixels and remains editable.
    pic=s.shapes.add_picture(str(path),Inches(x),Inches(y),width=Inches(w),height=Inches(h))
    if region:
        iw,ih,l,t,r,b=region
        pic.crop_left=l/iw;pic.crop_top=t/ih;pic.crop_right=1-r/iw;pic.crop_bottom=1-b/ih
    return pic

def base(num,title,kicker,notes):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(PAPER)
    rect(s,.48,.43,.47,.035,GOLD);text(s,1.08,.30,11,.3,kicker.upper(),11,GREEN,True)
    text(s,.55,.88,12.2,.66,title,32,INK,False,True)
    rect(s,.55,7.02,12.23,.012,'CCD0BA')
    text(s,.55,7.15,11.6,.2,'GOAL 007 / LOCAL COMPONENTS / CURRENT PRACTICE EVIDENCE / HOSTED ACCEPTANCE PENDING',8,MUTED)
    text(s,12.33,7.12,.5,.26,f'{num:02}',11,GREEN,True)
    s.notes_slide.notes_text_frame.text=notes+'\nAvailability snapshot: public HTTPS and current Live acceptance pending; source Private pending historical disclosure. Capture UI source '+SOURCE+'.'
    return s

def point(s,y,label,body,x=9.57,w=3.0):
    text(s,x,y,w,.4,label,18,INK,True);text(s,x,y+.52,w,.9,body,16,MUTED)

# 01: native editable title beside an untouched crop of the original product illustration.
s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(INK)
picture(s,OUT/'cover-1920x1080.png',6.50,0,6.833333,7.5,(1920,1080,936,0,1920,1080))
text(s,.62,.65,5.5,.5,'COOPERATION THROUGH CONVERSATION',14,LIGHT,True)
text(s,.59,1.63,5.8,1.8,'Talk Me\nHome.',58,PAPER,False,True)
text(s,.63,4.12,5.4,1.2,'You hold the plans.\nPip has eyes and hands.',25,PAPER)
text(s,.64,5.65,5.0,.84,'Share clues. Revise your route.\nBring your partner home.',19,LIGHT)
text(s,.64,6.86,5.2,.32,'Local prototype / Goal 007 / release evidence pending',10,LIGHT)
s.notes_slide.notes_text_frame.text='Original product illustration, not recorded gameplay. Native title/body text is editable. The cropped image is from the unchanged Goal006 cover; SHA256 '+COVER_HASH+'. This is a local working deck pending hosted acceptance and source disclosure.'

# 02: imperfect shared knowledge and a visible player-authored plan.
s=base(2,'A route is a hypothesis.','The cooperative loop','Actual 1920x1080 deterministic Practice screenshot: Relay Gallery. Brown route marks are manual player annotations. They are not robot telemetry or automatically communicated knowledge. No real provider was contacted. Full screenshot is included separately.')
picture(s,OUT/'screenshots/practice-gallery.png',.55,1.90,8.6,4.8375)
point(s,1.93,'Read the plan','Your atlas shows connections. Pip reports what is nearby.')
point(s,3.38,'Compare the evidence','An open gate can still be blocked. Ask before committing.')
point(s,4.83,'Revise together','Choose another path and confirm the exact move.')
text(s,.70,6.79,8.3,.19,'CURRENT PRACTICE / player-marked route / no live location marker',9,GREEN,True)

# 03: the console proves exact owner confirmation, with Dock collaboration inset.
s=base(3,'Say it. Read it. Confirm it.','Conversation becomes a decision','Two actual current Practice screenshots, with native PowerPoint crops. Left: Cargo pending Engage the Latch proposal, not yet executed. Right: Return Dock energy Stored after a confirmed hold-contact action. Physical tools require exact owner-bound UI confirmation before server commit. Human Power/Relay/charge controls are separate. No hands-free claim.')
picture(s,OUT/'screenshots/practice-cargo.png',.55,1.90,4.57,4.19,(1920,1080,1139,173,1752,735))
text(s,5.45,1.94,6.75,.75,'Pip proposes a physical action.\nYou choose whether it happens.',23,INK,True)
text(s,5.45,3.02,6.69,.7,'A spoken yes is not a console confirmation. The server checks the current state again when you confirm.',18,MUTED)
picture(s,OUT/'screenshots/practice-dock.png',5.45,4.05,6.67,1.70,(1920,1080,168,172,1116,414))
text(s,5.45,5.98,6.75,.65,'At the Dock, Pip holds the contact while you charge and store energy. Returning still needs a separate grant.',17,INK)
text(s,.68,6.72,11.5,.19,'CURRENT PRACTICE / exact pending proposal at left / acknowledged Dock instruments at right',9,GREEN,True)

# 04: an optional objective has a visible, verified consequence.
s=base(4,'A small memory can come home too.','The optional flight recorder','Actual current Practice confirmed-home screenshot. The optional recorder was selected before starting, locally inspected, picked up only after the exact Secure the flight recorder confirmation, and displayed after verified home. Direct rescue is valid; no hidden recording or provider speech is invented. The arrival scene is an illustration, as labelled in the game.')
picture(s,OUT/'screenshots/practice-home.png',.55,1.90,8.6,4.8375)
point(s,1.93,'Choose the detour','The Leaf archive is optional. A direct rescue is complete too.')
point(s,3.38,'Ask about the case','A worn label and faded star give the object a local story.')
point(s,4.83,'Give it a place','After a confirmed pickup and return, the case appears beside Pip.')
text(s,.70,6.79,8.3,.19,'CURRENT PRACTICE / confirmed home / optional recorder secured',9,GREEN,True)

# 05: all architecture remains editable native shapes and text.
s=base(5,'A conversation. A guarded action.','How the system works','Conceptual architecture of the implemented integration, not new Live test evidence. AssemblyAI Voice Agent API uses browser audio and role-safe local function tools; server validates proposals and exact owner confirmation. Provider credentials stay on server. Human maps/annotations do not enter automatic robot context. Practice is deterministic and offline. No screenshot or flattened diagram on this slide.')
rect(s,.65,2.02,3.70,2.29,PALE,radius=True);text(s,.89,2.25,3.2,.42,'HUMAN / BROWSER',16,INK,True);text(s,.89,2.92,3.15,1.11,'Private documents\nMicrophone + raw captions\nControls + confirmation',18,INK)
rect(s,4.83,2.02,3.66,2.29,INK,radius=True);text(s,5.08,2.25,3.15,.42,'ASSEMBLYAI / PIP',16,PAPER,True);text(s,5.08,2.92,3.15,1.11,'Speech + conversation\nLocal tool requests\nReply audio to browser',18,PAPER)
rect(s,9.00,2.02,3.68,2.29,PALE,radius=True);text(s,9.24,2.25,3.18,.42,'GAME SERVER',16,INK,True);text(s,9.24,2.92,3.10,1.11,'Mission state authority\nProposal validation\nConfirmed physical commit',18,INK)
text(s,4.39,2.83,.45,.6,'↔',26,GREEN,True);text(s,8.54,2.83,.45,.6,'↔',26,GREEN,True)
rect(s,.65,4.74,12.03,.68,'E9DFC6');text(s,.88,4.94,11.58,.3,'Observe → share a clue → propose → exact human confirmation → validate → act',19,INK,True)
text(s,.87,5.76,5.37,.71,'Separate knowledge\nPrivate map marks stay with the human.',18,INK)
text(s,7.07,5.76,5.25,.71,'Separate authority\nConversation alone cannot commit an action.',18,INK)
text(s,.88,6.62,11.8,.22,'TypeScript · Node.js 24 · React · Vite · Web Audio · Playwright / tool requests are routed through the browser',12,MUTED)

# 06: exact current availability; never invent URLs or a new success.
s=base(6,'A working prototype, with a clear next gate.','Availability and evidence','This slide is intentionally a pending-status snapshot and must be updated only from actual hosted URLs, current candidate Live evidence, visibility receipt and final media verification. Historical Goal005 had one synthetic-microphone real Voice success with ending ACK and one same-build failure without ACK. Neither is proof for the current build. No URL or public-availability claim is invented.')
rect(s,.63,1.99,5.83,3.38,PALE,radius=True);text(s,.92,2.25,5.23,.5,'READY LOCALLY',15,GREEN,True)
text(s,.92,2.96,5.07,1.94,'Practice and Training\nRescue + optional recorder\nExact confirmation + recovery\nFour screenshots + editable deck',19,INK)
rect(s,6.83,1.99,5.87,3.38,INK,radius=True);text(s,7.12,2.25,5.22,.5,'PENDING RELEASE EVIDENCE',15,LIGHT,True)
text(s,7.12,2.96,5.13,1.94,'Public HTTPS access: pending\nCurrent Live acceptance: pending\nSource: Private; historical disclosure pending\nFinal current Voice video: pending',19,PAPER)
text(s,.91,5.74,11.44,.65,'An earlier build completed one real Voice rescue; another run stalled. That history remains preserved and is not a current-build acceptance result.',18,MUTED)
text(s,.91,6.57,11.4,.25,'Screens shown here are deterministic Practice. No new human playtest or enjoyment result is claimed.',12,GREEN,True)

assert len(prs.slides)==6
for i,s in enumerate(prs.slides,1):
    for sh in s.shapes:
        assert sh.left>=0 and sh.top>=0 and sh.left+sh.width<=prs.slide_width+2 and sh.top+sh.height<=prs.slide_height+2,(i,sh.name)
prs.save(OUT/'Talk_Me_Home_Goal007_Deck.pptx')
print(json.dumps({'slides':len(prs.slides),'pptx':str(OUT/'Talk_Me_Home_Goal007_Deck.pptx'),'coverSha256':COVER_HASH}))

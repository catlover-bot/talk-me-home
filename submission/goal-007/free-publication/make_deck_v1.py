"""Create a versioned Free Practice publication deck, native PDF, renders and receipt.

Run with the existing submission-tools Python environment. Native LibreOffice
exports the saved PPTX; Poppler validates and renders it. Historical recipes are read-only.
"""
from pathlib import Path
import hashlib, json, subprocess, tempfile
from datetime import datetime, timezone
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR

ROOT=Path(__file__).resolve().parents[3]
ASSETS=ROOT/'submission/goal-007'
OUT=ASSETS/'free-publication'
OUT.mkdir(parents=True,exist_ok=True)
HOSTED='caa896d3a90dd4e8cb26499dba586b646bb7030a'
BASENAME='Talk_Me_Home_Goal007_Free_Practice_v1'
for suffix in ['.pptx','.pdf']:assert not (OUT/(BASENAME+suffix)).exists(),'Preserve existing deck; choose a new version.'
assert not (OUT/'deck-verification-v1.json').exists(),'Preserve existing receipt.'
PAPER='F4EFE2'; INK='24382F'; GREEN='58745D'; GOLD='B58A46'; MUTED='687365'; PALE='E2E5CF'; LIGHT='DCE3D1'
SOURCE='66f1bd206057720275fc9d56a6c9fbe72e409bcd'
COVER_HASH='38aa0826a78448c03c99bf717d44817b5aa57b3133071ff05ff6bf0df740ec91'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
original=ROOT/'submission/goal-006/cover-1920x1080.png'
assert sha(original)==COVER_HASH
assert sha(ASSETS/'cover-1920x1080.png')==COVER_HASH
capture=json.loads((ASSETS/'screenshots/capture-provenance.json').read_text())
assert capture['uiSourceCommit']==SOURCE and capture['providerTokenRequests']==0
assert capture['confirmedHome'] and capture['recorderSecuredAtHome']
for item in capture['screenshots']: assert sha(ASSETS/item['path'])==item['sha256']
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
prs.core_properties.title='Talk Me Home - Free Practice publication v1'
prs.core_properties.subject='Verified public Practice at caa896d; Live Voice and Live Text unavailable'
prs.core_properties.author='Talk Me Home'
prs.core_properties.keywords='Goal007, Free, hosted Practice, no real Voice claim'

def text(s,x,y,w,h,value,size=20,color=INK,bold=False,serif=False):
    box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=box.text_frame
    tf.clear();tf.word_wrap=True;tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(value.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line
        p.font.name='Liberation Serif' if serif else 'Liberation Sans';p.font.size=Pt(size)
        p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(8)
    return box

def rect(s,x,y,w,h,fill,line=None,radius=False):
    box=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    box.fill.solid();box.fill.fore_color.rgb=RGBColor.from_string(fill)
    box.line.fill.background() if not line else None
    if line:box.line.color.rgb=RGBColor.from_string(line)
    return box

def picture(s,path,x,y,w,h,region=None):
    # Native PowerPoint crop preserves original pixels and remains editable.
    pic=s.shapes.add_picture(str(path),Inches(x),Inches(y),width=Inches(w),height=Inches(h))
    if region:
        iw,ih,l,t,r,b=region
        pic.crop_left=l/iw;pic.crop_top=t/ih;pic.crop_right=1-r/iw;pic.crop_bottom=1-b/ih
    return pic

def base(num,title,kicker,notes):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(PAPER)
    rect(s,.48,.43,.47,.035,GOLD);text(s,1.08,.30,11,.3,kicker.upper(),11,GREEN,True)
    text(s,.55,.88,12.2,.66,title,32,INK,False,True)
    rect(s,.55,7.02,12.23,.012,'CCD0BA')
    text(s,.55,7.15,11.6,.2,'GOAL 007 / FREE PRACTICE PUBLICATION V1 / 30 SEP 2026 / LIVE UNAVAILABLE',8,MUTED)
    text(s,12.33,7.12,.5,.26,f'{num:02}',11,GREEN,True)
    s.notes_slide.notes_text_frame.text=notes+'\nPublic Practice verified at '+HOSTED+'. Live Voice and Live Text unavailable. Older local capture source '+SOURCE+'.'
    return s

def point(s,y,label,body,x=9.57,w=3.0):
    text(s,x,y,w,.4,label,18,INK,True);text(s,x,y+.52,w,.9,body,16,MUTED)

# 01: native editable title beside an untouched crop of the original product illustration.
s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(INK)
picture(s,ASSETS/'cover-1920x1080.png',6.50,0,6.833333,7.5,(1920,1080,936,0,1920,1080))
text(s,.62,.65,5.5,.5,'COOPERATION THROUGH CONVERSATION',14,LIGHT,True)
text(s,.59,1.63,5.8,1.8,'Talk Me\nHome.',58,PAPER,False,True)
text(s,.63,4.12,5.4,1.2,'You hold the plans.\nPip has eyes and hands.',25,PAPER)
text(s,.64,5.65,5.0,.84,'Share clues. Revise your route.\nBring your partner home.',19,LIGHT)
text(s,.64,6.86,5.2,.32,'Public Practice / Goal 007 / Free hosting / 30 Sep 2026',10,LIGHT)
s.notes_slide.notes_text_frame.text='Original product illustration, not recorded gameplay. Native title/body text is editable. The cropped image is from the unchanged Goal006 cover; SHA256 '+COVER_HASH+'. This is a Free Practice publication snapshot. Public Practice is verified at caa896d; Voice is unavailable.'

# 02: imperfect shared knowledge and a visible player-authored plan.
s=base(2,'A route is a hypothesis.','The cooperative loop','Earlier local 1920x1080 deterministic Practice capture, source 66f1bd: Relay Gallery. Brown route marks are manual player annotations. They are not robot telemetry or automatically communicated knowledge. No real provider was contacted. Full screenshot is included separately.')
picture(s,ASSETS/'screenshots/practice-gallery.png',.55,1.90,8.6,4.8375)
point(s,1.93,'Read the plan','Your atlas shows connections. Pip reports what is nearby.')
point(s,3.38,'Compare the evidence','An open gate can still be blocked. Ask before committing.')
point(s,4.83,'Revise together','Choose another path and confirm the exact move.')
text(s,.70,6.79,8.3,.19,'EARLIER LOCAL PRACTICE / source 66f1bd / player-marked route / no live location marker',9,GREEN,True)

# 03: the console proves exact owner confirmation, with Dock collaboration inset.
s=base(3,'Say it. Read it. Confirm it.','Conversation becomes a decision','Two earlier local Practice screenshots from source 66f1bd, with native PowerPoint crops. Left: Cargo pending Engage the Latch proposal, not yet executed. Right: Return Dock energy Stored after a confirmed hold-contact action. Physical tools require exact owner-bound UI confirmation before server commit. Human Power/Relay/charge controls are separate. No hands-free claim.')
picture(s,ASSETS/'screenshots/practice-cargo.png',.55,1.90,4.57,4.19,(1920,1080,1139,173,1752,735))
text(s,5.45,1.94,6.75,.75,'Pip proposes a physical action.\nYou choose whether it happens.',23,INK,True)
text(s,5.45,3.02,6.69,.7,'A spoken yes is not a console confirmation. The server checks the current state again when you confirm.',18,MUTED)
picture(s,ASSETS/'screenshots/practice-dock.png',5.45,4.05,6.67,1.70,(1920,1080,168,172,1116,414))
text(s,5.45,5.98,6.75,.65,'At the Dock, Pip holds the contact while you charge and store energy. Returning still needs a separate grant.',17,INK)
text(s,.68,6.72,11.5,.19,'EARLIER LOCAL PRACTICE / source 66f1bd / exact proposal and acknowledged Dock instruments',9,GREEN,True)

# 04: an optional objective has a visible, verified consequence.
s=base(4,'A small memory can come home too.','The optional flight recorder','Current hosted Practice confirmed-home screenshot at caa896d. The optional recorder was selected before starting, locally inspected, picked up only after the exact Secure the flight recorder confirmation, and displayed after verified home. Direct rescue is valid; no hidden recording or provider speech is invented. The arrival scene is an illustration, as labelled in the game.')
picture(s,ROOT/'artifacts/goal-007/hosted-free/collected-home.png',.55,1.90,7.74,4.8375)
point(s,1.93,'Choose the detour','The Leaf archive is optional. A direct rescue is complete too.')
point(s,3.38,'Ask about the case','A worn label and faded star give the object a local story.')
point(s,4.83,'Give it a place','After a confirmed pickup and return, the case appears beside Pip.')
text(s,.70,6.79,8.3,.19,'CURRENT HOSTED PRACTICE / caa896d / confirmed home / recorder secured',9,GREEN,True)

# 05: the active published path, with implemented Voice explicitly distinguished.
s=base(5,'A conversation. A guarded action.','How public Practice works','Native editable architecture. Public Practice runs the deterministic Pip simulator in the browser and the authoritative game server. The implemented AssemblyAI Voice Agent API integration is disabled on this public host. This slide is not evidence of a current real Voice connection. Private documents and map marks stay with the human. Every physical action still requires the exact owner confirmation.')
rect(s,.65,2.02,3.70,2.22,PALE,radius=True);text(s,.89,2.25,3.2,.42,'HUMAN / BROWSER',16,INK,True);text(s,.89,2.90,3.15,1.11,'Private documents\nTyped conversation\nControls + confirmation',18,INK)
rect(s,4.83,2.02,3.66,2.22,INK,radius=True);text(s,5.08,2.25,3.15,.42,'PRACTICE / PIP',16,PAPER,True);text(s,5.08,2.90,3.15,1.11,'Deterministic simulator\nLocal tool requests\nFaithful text captions',18,PAPER)
rect(s,9.00,2.02,3.68,2.22,PALE,radius=True);text(s,9.24,2.25,3.18,.42,'GAME SERVER',16,INK,True);text(s,9.24,2.90,3.10,1.11,'Mission state authority\nProposal validation\nConfirmed physical commit',18,INK)
text(s,4.39,2.83,.45,.6,'↔',26,GREEN,True);text(s,8.54,2.83,.45,.6,'↔',26,GREEN,True)
rect(s,.65,4.60,12.03,.68,'E9DFC6');text(s,.88,4.80,11.58,.3,'Observe → share a clue → propose → exact human confirmation → validate → act',19,INK,True)
text(s,.87,5.66,11.6,.43,'Implemented Voice integration — disabled on the public demo',19,INK,True)
text(s,.87,6.17,11.6,.60,'AssemblyAI Voice Agent API + browser audio and raw captions.\nThe published Practice path makes no provider calls; private map marks stay with the human.',16,MUTED)

# 06: current observed access, URLs and limitations; no current Voice claim.
s=base(6,'Try Practice in your browser.','Published access and its limits','The existing public origin and repository URL come from the supplied safe hosted verification receipt. HTTP and fresh-browser Practice checks passed at version 0.7.0 / caa896d. Both recorder-collected and selected-skipped rescues reached confirmed home. Free hosting and no persistent disk are owner-reported, not an authenticated dashboard audit. Live Voice and Live Text are disabled. Free hosting may cold-start; mission state is in memory and is lost when the server restarts. Platform selection: Other. No form submission or media upload has been performed.')
text(s,.67,1.91,5.70,.36,'PUBLIC DEMO',13,GREEN,True)
a=text(s,.67,2.38,5.9,.49,'https://talk-me-home.onrender.com',20,INK,True)
a.text_frame.paragraphs[0].runs[0].hyperlink.address='https://talk-me-home.onrender.com'
text(s,.67,3.08,5.70,.33,'REPOSITORY',13,GREEN,True)
a=text(s,.67,3.51,5.95,.45,'github.com/catlover-bot/talk-me-home',17,INK,True)
a.text_frame.paragraphs[0].runs[0].hyperlink.address='https://github.com/catlover-bot/talk-me-home'
text(s,.67,4.26,5.73,.9,'Verified: Cargo → Gallery → Dock → home.\nRecorder collected and skipped outcomes.\nExact confirmations, Pause/resume and replay.',16,INK)
text(s,.67,5.50,5.70,.30,'Platform: Other · v0.7.0 · hosted build caa896d',12,GREEN,True)
picture(s,ROOT/'artifacts/goal-007/hosted-free/entrance-1280.png',6.91,1.98,5.77,3.245625)
text(s,7.03,5.30,5.53,.24,'CURRENT HOSTED STILL / caa896d / Practice available',10,GREEN,True)
rect(s,.65,5.86,12.03,.92,INK,radius=True)
text(s,.87,6.01,11.6,.30,'PRACTICE ONLY · LIVE VOICE & LIVE TEXT UNAVAILABLE',17,PAPER,True)
text(s,.87,6.42,11.6,.22,'Free hosting may cold-start. Server restarts lose mission progress. Owner reports Free / no persistent disk.',12,LIGHT)

assert len(prs.slides)==6
for i,s in enumerate(prs.slides,1):
    for sh in s.shapes:
        assert sh.left>=0 and sh.top>=0 and sh.left+sh.width<=prs.slide_width+2 and sh.top+sh.height<=prs.slide_height+2,(i,sh.name)

# Snapshot existing historical deliverables; all new versioned outputs stay here.
preserved={str(p.relative_to(ROOT)):sha(p) for p in (ROOT/'submission').rglob('*') if p.is_file() and OUT not in p.parents and (p.suffix.lower() in {'.pptx','.pdf','.mp4','.zip'} or 'manifest' in p.name)}
hosted=json.loads((ROOT/'artifacts/goal-007/free-hosted-verification.json').read_text())
assert hosted['status']=='HOSTED_PRACTICE_PASS' and hosted['observedIdentity']['commit']==HOSTED
assert hosted['realVoiceVerification']=='NOT_RUN_LIVE_DISABLED_BY_CURRENT_SCOPE'
for item in hosted['screenshots']:
    assert sha(ROOT/item['path'])==item['sha256']
prs.save(OUT/(BASENAME+'.pptx'))
with tempfile.TemporaryDirectory(prefix='tmh-free-deck-lo-') as profile:
    result=subprocess.run(['libreoffice','-env:UserInstallation='+Path(profile).as_uri(),'--headless','--convert-to','pdf','--outdir',str(OUT),str(OUT/(BASENAME+'.pptx'))],capture_output=True,text=True,check=True,timeout=60)
    assert (OUT/(BASENAME+'.pdf')).exists(),result.stdout+result.stderr
renders=OUT/'renders';renders.mkdir(exist_ok=True)
subprocess.run(['pdftoppm','-png','-scale-to-x','1920','-scale-to-y','1080',str(OUT/(BASENAME+'.pdf')),str(renders/'slide')],check=True,timeout=60)
for index in range(1,7):
    old=renders/f'slide-{index}.png';new=renders/f'slide-{index:02}.png'
    if old.exists():old.rename(new)
from PIL import Image
pdfinfo=subprocess.check_output(['pdfinfo',str(OUT/(BASENAME+'.pdf'))],text=True)
assert 'Pages:           6' in pdfinfo and 'LibreOffice' in pdfinfo
text_by_page=subprocess.check_output(['pdftotext','-layout',str(OUT/(BASENAME+'.pdf')),'-'],text=True).split('\f')
assert 'https://talk-me-home.onrender.com' in text_by_page[5]
assert 'caa896d' in text_by_page[5] and '66f1bd' in text_by_page[1] and '66f1bd' in text_by_page[2]
assert 'LIVE VOICE & LIVE TEXT UNAVAILABLE' in text_by_page[5]
assert all(Image.open(renders/f'slide-{i:02}.png').size==(1920,1080) for i in range(1,7))
assert all(sha(ROOT/path)==digest for path,digest in preserved.items()),'Historical deliverable changed'
outputs=[OUT/(BASENAME+'.pptx'),OUT/(BASENAME+'.pdf')]+[renders/f'slide-{i:02}.png' for i in range(1,7)]
receipt={
 'schemaVersion':1,'status':'FREE_PRACTICE_PUBLICATION_DECK_V1','createdAt':datetime.now(timezone.utc).isoformat(),
 'scope':'New local versioned six-slide deck/PDF only; no app, provider, upload or deployment request',
 'hostedBuild':hosted['observedIdentity'],'hostedReceipt':'artifacts/goal-007/free-hosted-verification.json',
 'earlierLocalGameplaySource':SOURCE,'hostedStillsSource':HOSTED,'realVoiceClaim':False,
 'providerRequests':0,'publicAppRequests':0,'coverSha256':COVER_HASH,
 'editableSlides':len(prs.slides),'pdfRenderer':'Native LibreOffice','renderSize':[1920,1080],
 'nativeTextShapes':[sum(1 for sh in s.shapes if sh.has_text_frame) for s in prs.slides],
 'nativePictureShapes':[sum(1 for sh in s.shapes if sh.shape_type==13) for s in prs.slides],
 'checks':['6 PDF pages','original capture hashes','required source and limitation labels','actual URLs','all shapes inside canvas','1920x1080 renders','historical deliverable hashes unchanged'],
 'visualInspection':'PENDING_ACTUAL_PAGE_REVIEW',
 'outputs':[{ 'path':str(p.relative_to(ROOT)), 'bytes':p.stat().st_size,'sha256':sha(p)} for p in outputs],
 'preservedHistoricalDeliverables':preserved,
 'inputScreenshots':[{ 'path':str(p.relative_to(ROOT)), 'sha256':sha(p)} for p in [ASSETS/'screenshots/practice-gallery.png',ASSETS/'screenshots/practice-cargo.png',ASSETS/'screenshots/practice-dock.png',ROOT/'artifacts/goal-007/hosted-free/collected-home.png',ROOT/'artifacts/goal-007/hosted-free/entrance-1280.png']]
}
(OUT/'deck-verification-v1.json').write_text(json.dumps(receipt,indent=2)+'\n')
print(json.dumps({'outputs':receipt['outputs'],'preserved':len(preserved),'slides':6},indent=2))

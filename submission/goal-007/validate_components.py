"""Write bounded English form copy and validate the local component snapshot."""
from pathlib import Path
from datetime import datetime, timezone
import hashlib, json, re, subprocess
from PIL import Image
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'submission/goal-007';WORK=ROOT/'.validation/goal-007-media'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
def save(path,value):path.write_text(json.dumps(value,indent=2,ensure_ascii=False)+'\n')
form={
 'status':'WORKING_LOCAL_COMPONENTS_PENDING_RELEASE_EVIDENCE',
 'title':'Talk Me Home',
 'summary':'You hold the plans. Pip has eyes and hands. Share clues, revise a route and confirm each action to bring your AI partner home, with an optional memory to carry back.',
 'description':'''Talk Me Home is a cooperative puzzle game about two incomplete views of one station. At Mission Control, you hold private documents and remote equipment controls. Pip, your AI partner, can observe the space nearby, inspect objects and propose physical actions. Conversation turns separate clues into a shared plan.

The rescue crosses Cargo Bay, Relay Gallery and Return Dock. A shared power supply demands coordination. A gate on the map may open onto blocked cargo, so Pip's report can change your route. At the Dock, one partner holds a contact while the other stores energy. Every physical proposal requires an exact human console confirmation before the server validates and commits it.

An optional flight recorder gives the journey a personal choice. Read the archive clue, ask Pip about the worn case and decide whether to make the detour. A confirmed pickup and return place the recorder beside Pip at home. A direct rescue is complete too.

Voice mode connects browser audio to the AssemblyAI Voice Agent API. The browser presents raw captions and local playback; the server owns mission state and action validation. Private map marks stay with the human. Practice provides a deterministic local introduction.

This package shows the current interface in Practice. Public HTTPS access, current-build Live acceptance and final Voice media remain pending. The repository remains Private pending historical disclosure. An earlier build completed one real Voice rescue and failed another run; that history is preserved separately.''',
 'additionalInformation':'''Local components include six editable slides, their native LibreOffice-rendered PDF, four current 1920x1080 Practice screenshots and the unchanged original cover. No Practice frame is presented as a new real Voice pass.

The optional recorder is selected before starting; its pickup uses the same exact confirmation boundary as every other physical action. Recovery requests ask for bounded observations and retain raw transcripts. Ambient radio sound is opt-in and gives way to actual microphone input or local reply playback.

Historical real-provider evidence used synthetic microphone input and digital playback, not a human speech or physical speaker test. One earlier rescue reached home with an ending ACK; a second run on the same old candidate stalled without an ACK. Those originals, failures and accounting remain preserved.

Public HTTPS, protected Live acceptance, source disclosure and the final current Voice video are still pending. There is no uploaded media or event-submission result in this local snapshot. URLs will be filled only after actual availability is verified.''',
 'technologies':'AssemblyAI Voice Agent API; TypeScript; Node.js 24; React; Vite; Web Audio; authoritative in-memory server; Playwright.',
 'demoApplicationPlatform':'Other',
 'demoApplicationPlatformDetail':'Browser-based web game',
 'demoUrl':None,'repositoryUrl':None,'videoUrl':None,'pitchDeckUrl':None,
 'availability':{'publicHttps':'pending','currentLiveAcceptance':'pending','source':'Private pending historical disclosure','finalCurrentVoiceMedia':'pending','eventSubmission':'not performed'}
}
counts={key:{'characters':len(form[key]),'words':len(form[key].split())} for key in ['title','summary','description','additionalInformation']}
assert 5<=counts['title']['characters']<=50
assert 50<=counts['summary']['characters']<=255
assert 600<=counts['description']['characters']<=2000 and counts['description']['words']>=100
assert counts['additionalInformation']['characters']<=2000
save(OUT/'submission-form.json',form)
md='# Talk Me Home - Goal 007 working form copy\n\nStatus: **local components; release evidence pending**. This is not a submitted form.\n\n'
for name,key in [('Title','title'),('Summary','summary'),('Description','description'),('Additional information','additionalInformation'),('Technologies','technologies')]:md+=f'## {name}\n\n{form[key]}\n\n'
md+='## Application and links\n\n- Demo Application Platform: Other (browser-based web game).\n- Public HTTPS demo: pending; no invented URL.\n- Public repository: Private pending historical disclosure.\n- Video URL: pending current Voice evidence and media preparation.\n- Pitch deck URL: not uploaded; local editable PPTX and native PDF included.\n\n'
md+='## Length checks\n\n| Field | Characters | Words |\n| --- | ---: | ---: |\n'+''.join(f'| {k} | {v["characters"]} | {v["words"]} |\n' for k,v in counts.items())
(OUT/'submission-form.md').write_text(md)

capture=json.loads((OUT/'screenshots/capture-provenance.json').read_text())
assert capture['providerTokenRequests']==0 and capture['confirmedHome'] and capture['recorderSecuredAtHome']
private=json.loads((WORK/'practice-source-private.json').read_text());raw=Path(private['path'])
assert sha(raw)==capture['video']['sha256']==private['sha256']
probe=json.loads(subprocess.check_output(['ffprobe','-v','error','-show_format','-show_streams','-of','json',str(raw)]))
streams=probe['streams'];assert len(streams)==1 and streams[0]['codec_type']=='video'
assert (streams[0]['width'],streams[0]['height'])==(1920,1080)
private['probe']=probe;save(WORK/'practice-source-private.json',private)
recipe={
 'schemaVersion':1,'status':'REUSABLE_EDIT_RECIPE_NOT_A_FINAL_VIDEO','target':{'durationSecondsRange':[180,240],'width':1920,'height':1080,'codec':'H.264','audio':'AAC','frameRate':30},
 'currentVoiceSource':None,
 'voiceSourceRequirements':['One preserved original current-candidate synthetic-microphone Voice run','Same connection across Cargo, Gallery, Dock and confirmed home','Exact UI confirmations visible','Actual provider ending ACK retained','Raw input, provider output, video and faithful transcript retained'],
 'practiceSource':{**capture['video'],'mode':'Practice - deterministic simulation','uiSourceCommit':capture['uiSourceCommit'],'durationSeconds':float(probe['format']['duration']),'audioStreams':0,'moments':capture['moments'],'allowedUse':'Separately labelled current Practice inserts; silent source, never provider speech'},
 'timelineOutline':[
  {'section':'Premise','plannedSeconds':10,'content':'Native title card and an actual current-source greeting when available'},
  {'section':'Cargo cooperation','plannedSeconds':40,'content':'Clue, short conversation, exact Latch confirmation, shared Power decision'},
  {'section':'Gallery revision','plannedSeconds':60,'content':'Actual report, uncertainty, changed plan and exact confirmation'},
  {'section':'Dock return','plannedSeconds':45,'content':'Contact held, charge stored, release/board, return grant and confirmed home'},
  {'section':'Optional recorder','plannedSeconds':15,'content':'Only actual current Voice or separately labelled Practice material'},
  {'section':'Technology and availability','plannedSeconds':20,'content':'Implemented architecture, verified actual access and evidence limits'}
 ],
 'editingRules':['Keep every quote verbatim and aligned to its own audio','Do not place old provider speech under current Practice screenshots','Label shortened waits and non-contiguous cuts','Preserve an intelligible cooperation exchange rather than a montage of confirmations','Use separately identified Practice inserts only when needed','Do not invent a successful end or ending ACK','Review final MP4 by listening and watching, then probe dimensions/codecs/duration','Update deck, form and availability from actual hosted outcomes before final packaging'],
 'historicalPolicy':'Goal005 originals and Goal006 recipe/binaries remain read-only; no historical clip is relabelled as current Live.'
}
save(OUT/'video-edit-recipe.json',recipe)

# Check original deliverables against their already recorded hashes, without rewriting any bytes.
historical_root=ROOT.parent/'talk-me-home-goal-006/submission/goal-006'
expected={
 'cover-1920x1080.png':'38aa0826a78448c03c99bf717d44817b5aa57b3133071ff05ff6bf0df740ec91',
 'Talk_Me_Home_Submission.mp4':'daa57ee6d785ca985a428028db461462df0d50bd1d753fd024fb2600d3f949dd',
 'Talk_Me_Home_Pitch.pptx':'2cd226b28d01fd6ae12eac0e5b877ec67d82d10bdfec6825ebd0d0d985417f1a',
 'Talk_Me_Home_Pitch.pdf':'ecb30f0cb53da8f8dfab1d3e0a54021b1a4e1f5991b6e47d8721aada2c90db5f',
 'Talk_Me_Home_Submission_Package.zip':'faf2705d8c6c763636c094657b4bb987fb51bb3e897cbbfe68d2cac83fb7f9ff'}
preserved=[]
for name,value in expected.items():
 path=historical_root/name;assert sha(path)==value
 preserved.append({'sourceId':'Goal006/'+name,'bytes':path.stat().st_size,'sha256':value})
for name in ['make_media.py','inspect_and_package.py','edit-plan.json','media-manifest.json']:
 path=historical_root/name;assert sha(path)==sha(ROOT/'submission/goal-006'/name)
 preserved.append({'sourceId':'Goal006/'+name,'bytes':path.stat().st_size,'sha256':sha(path),'comparedToUnmodifiedTrackedCopy':True})
oldvoice=ROOT.parent/'talk-me-home/.validation/goal-005-media/Goal005_Attempt12_Success_Uncut.mp4'
assert sha(oldvoice)=='d69a0b9a1e793294399e8731cffb7a4c8765772e685510598ff8b7d04e8d4812'
preserved.append({'sourceId':'Goal005/Attempt12/SuccessUncut','bytes':oldvoice.stat().st_size,'sha256':sha(oldvoice)})

pptx=Presentation(OUT/'Talk_Me_Home_Goal007_Deck.pptx');assert len(pptx.slides)==6
slideStats=[]
for i,slide in enumerate(pptx.slides,1):
 count=sum(1 for shape in slide.shapes if shape.has_text_frame and shape.text_frame.text.strip())
 assert count>=4
 slideStats.append({'page':i,'editableTextShapes':count,'totalShapes':len(slide.shapes),'speakerNotesPresent':bool(slide.notes_slide.notes_text_frame.text.strip())})
pdfinfo=subprocess.check_output(['pdfinfo',str(OUT/'Talk_Me_Home_Goal007_Deck.pdf')],text=True)
assert re.search(r'Pages:\s+6\b',pdfinfo) and 'LibreOffice' in pdfinfo
pdftext=subprocess.check_output(['pdftotext','-layout',str(OUT/'Talk_Me_Home_Goal007_Deck.pdf'),'-'],text=True)
for phrase in ['A route is a hypothesis.','Say it. Read it. Confirm it.','A small memory can come home too.','A conversation. A guarded action.','Public HTTPS access: pending','Current Live acceptance: pending','Source: Private; historical disclosure']: assert phrase in pdftext
files=[]
for path in sorted(OUT.rglob('*')):
 if path.is_file() and path.name!='media-components-manifest.json' and '__pycache__' not in path.parts:
  item={'path':path.relative_to(OUT).as_posix(),'bytes':path.stat().st_size,'sha256':sha(path)}
  if path.suffix=='.png':
   with Image.open(path) as im:item['width'],item['height']=im.size
  files.append(item)
manifest={
 'schemaVersion':1,'status':'WORKING_LOCAL_COMPONENTS_PENDING_RELEASE_EVIDENCE','generatedAtUtc':datetime.now(timezone.utc).isoformat(),
 'sourceCommit':capture['sourceCommit'],'uiSourceCommit':capture['uiSourceCommit'],'providerTokenRequests':0,
 'formChecks':counts,'pptxSlides':slideStats,'pdf':{'pages':6,'renderer':'Native LibreOffice Impress PDF export','visualReview':'Each of six rendered pages inspected; corrected architecture text overflow, availability spacing and Cargo crop'},
 'screenshots':{'count':4,'mode':'Practice - deterministic simulation','dimensions':'1920x1080','chapters':['Cargo Bay','Relay Gallery','Return Dock','Confirmed home with recorder'],'inspection':'All four originals visually inspected'},
 'cover':{'unchangedFrom':'Goal006/cover-1920x1080.png','sha256':expected['cover-1920x1080.png']},
 'practiceRecording':recipe['practiceSource'],'historicalPreservation':preserved,'files':files,
 'notCompleted':['Final current Voice video','Hosted HTTPS acceptance','Current-candidate Live acceptance','Historical source disclosure and public visibility verification','Uploaded media','Event submission'],
 'limits':['Practice capture is silent and is not real Voice evidence','This snapshot predates final hosted integration; screenshots must be compared with the final integrated UI before release','No current human playtest or measured enjoyment claim']
}
save(OUT/'media-components-manifest.json',manifest)
(WORK/'pdfinfo.txt').write_text(pdfinfo);(WORK/'pdf-text.txt').write_text(pdftext)
print(json.dumps({'formCounts':counts,'slides':slideStats,'files':[f for f in files if f['path'].endswith(('.pptx','.pdf'))],'practiceDuration':recipe['practiceSource']['durationSeconds'],'preserved':len(preserved)},indent=2))

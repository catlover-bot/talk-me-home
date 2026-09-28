"""Create a small local-deliverable receipt after rendering and visual inspection."""
from pathlib import Path
import argparse,hashlib,json,subprocess
from pptx import Presentation
from PIL import Image

ROOT=Path(__file__).resolve().parents[2]
parser=argparse.ArgumentParser(description='Record checks only after all final native slide renders and video transitions have been visually reviewed.')
parser.add_argument('--visual-review-complete',action='store_true',required=True)
parser.parse_args()
WORK=ROOT/'.validation/goal-004e-follow-up-media'
DEST=ROOT/'submission/local-deliverables'
def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def metadata(path):return {'path':str(path.relative_to(ROOT)),'bytes':path.stat().st_size,'sha256':sha(path)}
def command(args):return subprocess.check_output(args,cwd=ROOT,text=True,stderr=subprocess.STDOUT)
record=json.loads((WORK/'practice-recording.json').read_text())
timeline=json.loads((WORK/'edit-timeline.json').read_text())
pptx=ROOT/'submission/Talk_Me_Home_Pitch.pptx';pdf=ROOT/'submission/Talk_Me_Home_Pitch.pdf';video=ROOT/'submission/Talk_Me_Home_Demo_Draft.mp4'
deck=Presentation(pptx)
assert len(deck.slides)==6 and abs(deck.slide_width/deck.slide_height-16/9)<.00001
slides=[]
for n,slide in enumerate(deck.slides,1):
    editable=sum(shape.has_text_frame and bool(shape.text.strip()) for shape in slide.shapes)
    assert editable>=4
    render=WORK/'rendered-slides'/f'slide-{n}.png'
    assert Image.open(render).size==(1280,720)
    slides.append({'slide':n,'editableTextShapes':editable,'pictureShapes':sum(shape.shape_type==13 for shape in slide.shapes),'sourceNotesPresent':bool(slide.notes_slide.notes_text_frame.text.strip()),'nativeRender':metadata(render)})
media=json.loads(command(['ffprobe','-v','error','-show_entries','format=duration,size:stream=codec_name,codec_type,width,height,sample_rate,channels','-of','json',str(video)]))
assert 120<=float(media['format']['duration'])<=180
assert any(s['codec_type']=='video' and s['codec_name']=='h264' and s['width']==1280 and s['height']==720 for s in media['streams'])
assert any(s['codec_type']=='audio' and s['codec_name']=='aac' for s in media['streams'])
original=json.loads((ROOT/'artifacts/goal-004e/live-media.json').read_text())
original_checks=[]
for item in original['files']:
    path=ROOT/item['path']; matches=sha(path)==item['sha256'] and path.stat().st_size==item['bytes'];assert matches
    original_checks.append({'path':item['path'],'sha256':item['sha256'],'unchanged':matches})
receipt={'status':'LOCAL_DELIVERABLES_RENDERED','releaseStatus':'RELEASE_NOT_LIVE_VERIFIED','practice':{key:record[key] for key in ['commit','runtimeSha256','recordedAt','browser','providerRequests','externalRequests','webSockets','completed','galleryRoute','cleanup']},
 'practiceConfirmations':len(record['confirmations']),'outputs':[metadata(pptx),metadata(pdf),metadata(video)],'slides':slides,'video':media,
 'toolchain':{'pythonPptx':'1.0.2','pillow':'12.3.0','libreoffice':command(['libreoffice','--version']).strip(),'pdftoppm':command(['pdftoppm','-v']).splitlines()[0],'ffmpeg':command(['ffmpeg','-version']).splitlines()[0]},
 'historicalMediaPreservation':original_checks,'sourceModeBoundary':timeline['audioBoundary'],
 'inspection':{'nativePowerPointRenderedWith':'LibreOffice Impress PDF export; every slide then rasterized by Poppler for visual review','allSixSlidesInspected':True,'repairedIssues':['Evidence slide body wrapping/overflow corrected before final render.'],'videoFrames':'Representative frames at every mode/chapter transition and the final availability card inspected; labels persist in footage.','physicalDevicesOrHumanEnjoyment':'Not tested; no claim.'},
 'localRawPractice':metadata(WORK/'practice-source.webm'),'links':{'publicDemo':None,'uploadedVideo':None,'uploadedPresentation':None,'eventSubmission':None}}
DEST.mkdir(parents=True,exist_ok=True)
(DEST/'provenance.json').write_text(json.dumps(receipt,indent=2)+'\n')
(DEST/'edit-timeline.json').write_text(json.dumps(timeline,indent=2)+'\n')
print(json.dumps({'outputs':receipt['outputs'],'durationSeconds':media['format']['duration'],'slides':len(slides)},indent=2))

"""Build the editable local deck. No network, provider, or application mutation."""
from pathlib import Path
import argparse
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR

ROOT = Path(__file__).resolve().parents[2]
parser = argparse.ArgumentParser()
parser.add_argument('--captures', default='submission/local-deliverables/screenshots')
parser.add_argument('--output', default='submission/Talk_Me_Home_Pitch.pptx')
args = parser.parse_args()
CAP = ROOT / args.captures
prs = Presentation()
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)
prs.core_properties.title = 'Talk Me Home — voice-led co-op with explicit confirmation'
prs.core_properties.subject = 'Local submission draft; full Live Rescue remains unverified'
prs.core_properties.author = 'Talk Me Home project'
PAPER='F4EFE2'; INK='24382F'; GREEN='557862'; GOLD='B58A46'; MUTED='607066'; PALE='E4E9DD'; WHITE='FFFFFF'

def rgb(c): return RGBColor.from_string(c)
def rect(s,x,y,w,h,fill,line=None,radius=False):
    shape=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb=rgb(fill)
    if line: shape.line.color.rgb=rgb(line)
    else: shape.line.fill.background()
    if radius: shape.adjustments[0]=.08
    return shape
def text(s,copy,x,y,w,h,size=22,color=INK,bold=False,font='Liberation Sans',align=None):
    box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=box.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.TOP
    for i,line in enumerate(copy.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line; p.font.name=font; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=rgb(color)
        p.space_after=Pt(7); p.line_spacing=1.08
        if align is not None: p.alignment=align
    return box
def picture(s,path,x,y,w,h,crop=False):
    from PIL import Image
    path=Path(path); iw,ih=Image.open(path).size
    if crop:
        pic=s.shapes.add_picture(str(path),Inches(x),Inches(y),width=Inches(w),height=Inches(h))
        ratio=iw/ih; target=w/h
        if ratio>target: pic.crop_left=pic.crop_right=(1-target/ratio)/2
        else: pic.crop_top=pic.crop_bottom=(1-ratio/target)/2
    else:
        scale=min(w/iw,h/ih); pw=iw*scale; ph=ih*scale
        pic=s.shapes.add_picture(str(path),Inches(x+(w-pw)/2),Inches(y+(h-ph)/2),width=Inches(pw),height=Inches(ph))
    return pic
def slide(number,kicker,title,subtitle=None):
    s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb(PAPER)
    rect(s,.65,.49,.42,.055,GOLD)
    text(s,kicker.upper(),1.2,.39,10,.35,11,MUTED,True)
    text(s,title,.65,.94,12,1.0,36,INK,False,'Liberation Serif')
    if subtitle: text(s,subtitle,.68,1.85,11.9,.55,17,MUTED)
    text(s,'TALK ME HOME',.68,7.04,7,.22,9,MUTED,True)
    text(s,f'{number:02d} / 06',11.75,7.02,.9,.25,10,MUTED)
    return s
def notes(s,copy): s.notes_slide.notes_text_frame.text=copy

s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb(INK)
art=picture(s,ROOT/'submission/assets/cover-illustration-1600x900.png',7.1,0,6.24,7.5)
# Keep the original building and Pip portion of the cover; the title is native editable text.
art.left=Inches(7.1); art.top=0; art.width=Inches(6.24); art.height=Inches(7.5)
art.crop_left=.532; art.crop_right=0; art.crop_top=0; art.crop_bottom=0
rect(s,0,0,7.15,7.5,INK)
text(s,'A VOICE-LED CO-OP RESCUE GAME',.75,.74,6,.45,12,'CBD3BC',True)
text(s,'Talk Me\nHome.',.73,1.47,6.2,2.15,59,'F4EEDC',False,'Liberation Serif')
text(s,'You have the map.\nPip has eyes and hands.',.78,4.1,5.7,1.0,24,'EADBBC')
text(s,'Share clues. Choose a plan.\nConfirm each proposed action on the console.',.78,5.65,5.9,.95,17,'DCE2CE')
text(s,'Original game artwork · local submission draft',.78,7.04,6,.25,10,'CBD3BC')
notes(s,'Source: submission/project-description.md; existing original cover illustration. Artwork is not gameplay evidence. Native title and body text remain editable. Contract: spoken/text coordination plus explicit owner UI confirmation, never hands-free permission inference.')

s=slide(2,'The cooperation premise','Neither partner has the whole picture.','Conversation connects a reference map with a local point of view.')
rect(s,.68,2.68,5.75,3.4,'E8E1CF',radius=True); rect(s,6.89,2.68,5.75,3.4,PALE,radius=True)
text(s,'MISSION CONTROL',.98,3.02,5,.42,15,GREEN,True)
text(s,'Read the documents.\nControl remote equipment.\nConfirm one exact proposal.',.98,3.73,5.1,1.8,23)
text(s,'Private notes stay at your desk.',.98,5.64,5,.4,14,MUTED)
text(s,'PIP',7.2,3.02,5,.42,15,GREEN,True)
text(s,'Observe nearby equipment.\nInspect without permission.\nPropose a local action.',7.2,3.73,5.05,1.8,23)
text(s,'No automatic view of your map or notes.',7.2,5.64,5,.4,14,MUTED)
text(s,'A proposal communicates intent. Only a confirmed, server-validated action changes the world.',.8,6.43,11.9,.4,16,GREEN)
notes(s,'Sources: game/shared/contracts.ts, game/server/proposals.ts and sessions.ts; current project description. The human does not receive a live local robot survey; Pip does not receive the human manual, atlas, notes or hidden solution. This is a normal gameplay information split, not a browser developer-tools security claim.')

s=slide(3,'Three chapters','One rescue. Three ways to cooperate.','New deterministic Practice captures · the same explicit confirmation boundary')
for x,name,file,body,n in [(.68,'Cargo Bay','cargo.png','Reason about shared machinery\nand remote Power.','01'),(4.9,'Relay Gallery','gallery.png','Compare spoken clues with an atlas.\nRethink the route when blocked.','02'),(9.12,'Return Dock','dock.png','Coordinate charging and readiness.\nAuthorize the return separately.','03')]:
    rect(s,x,2.65,3.53,2.38,WHITE); picture(s,CAP/file,x+.02,2.67,3.49,2.34)
    text(s,n,x,5.26,.5,.35,13,GOLD,True); text(s,name,x+.43,5.16,3.2,.55,23,INK,False,'Liberation Serif')
    text(s,body,x,5.87,3.62,.9,15,MUTED)
notes(s,'Images: newly captured local compiled production Practice, from the follow-up build recorded in submission/local-deliverables/provenance.json. Deterministic typed Practice is offline and is not evidence of real model completion. Chapters and authored mechanics are unchanged.')

s=slide(4,'The interaction','Talk with Pip. Confirm on the console.')
rect(s,.68,2.15,8.1,4.56,WHITE); picture(s,CAP/'pending.png',.69,2.16,8.08,4.54)
for y,n,title,body in [(2.2,'1','Conversation','Share a clue or request.'),(3.32,'2','A specific proposal','Read the exact action.'),(4.44,'3','Human confirmation','Confirm, or choose Not yet.'),(5.56,'4','Verified result','The server rechecks conditions.')]:
    rect(s,9.13,y,.43,.43,GREEN,radius=True); text(s,n,9.27,y+.055,.22,.3,14,WHITE,True)
    text(s,title,9.79,y-.02,2.85,.46,19,INK,True); text(s,body,9.79,y+.48,2.85,.55,14,MUTED)
text(s,'Current Practice UI · spoken “yes” does not execute an action',.73,6.76,8.6,.28,11,MUTED)
notes(s,'Current UI screenshot is an actual local Practice screen, not a recreated mockup. Sources: ActionProposalStrip/useMission, SessionStore exact immutable proposals. Read-only calls do not require confirmation. Follow-up adds a bounded acknowledgement opportunity, without issuing a second result for the original tool call or auto-confirming another action.')

s=slide(5,'Implementation and evidence','An implemented game. A clearly scoped result.')
rect(s,.68,2.24,3.7,3.99,PALE,radius=True); rect(s,4.82,2.24,3.7,3.99,'E8E1CF',radius=True); rect(s,8.96,2.24,3.7,3.99,'ECE6DF',radius=True)
text(s,'BUILT',.98,2.6,3.15,.34,13,GREEN,True)
text(s,'React + TypeScript\nNode game server\nAssemblyAI Voice API',.98,3.18,3.15,1.8,20)
text(s,'Exact proposals, owner decisions\nand faithful captions.',.98,5.34,3.15,.72,14,MUTED)
text(s,'OBSERVED IN REAL VOICE',5.12,2.6,3.1,.34,12,GREEN,True)
text(s,'0 unconfirmed commits\n1 confirmed Latch action\nRemote end acknowledged',5.12,3.18,3.04,1.85,19)
text(s,'Synthetic voice + UI confirmation.\nStalled before the Cargo crossing.',5.12,5.34,3.08,.72,13,MUTED)
text(s,'VERIFIED OFFLINE',9.26,2.6,3.1,.34,13,GREEN,True)
text(s,'Both Gallery variants\nDock and confirmed home\nBounded recovery cases',9.26,3.18,3.05,1.8,20)
text(s,'Practice and explicitly fake peers.\nNot a real model pass.',9.26,5.34,3.05,.68,14,MUTED)
text(s,'Full Live Rescue verification remains pending.',.82,6.52,11.8,.44,22,INK,True)
notes(s,'Real evidence: artifacts/goal-004e/live/2026-09-28T10-24-13-651Z-voice-mission-conversation.md and live-media.json, execution bb6dfd2181503415a14c440c52dd762db88f5e3c. Three unconfirmed inputs produced zero physical change, then one exact Latch confirmation committed once; Rescue failed. One session.end/session.ended and zero cleanup survivors. Follow-up offline tests and build identity are in current validation/provenance; no new provider use. No physical microphone, loudspeaker, human natural completion or enjoyment result is claimed.')

s=slide(6,'Product direction and availability','A small robot. A shared way home.')
picture(s,CAP/'home.png',.68,2.22,6.83,3.85)
text(s,'PLAYABLE LOCALLY',8.02,2.36,4.5,.35,13,GREEN,True)
text(s,'Three-chapter Rescue\nand optional Training.',8.02,2.94,4.57,.98,25)
text(s,'Next: evaluate natural play\nand explore more authored episodes.',8.02,4.27,4.55,.82,18,MUTED)
text(s,'Public demo URL: not deployed.\nFull Live acceptance: pending.\nUpload and event entry: not submitted.',8.02,5.49,4.61,1.04,16,INK)
text(s,'Practice ending shown · a product direction, not a claim of customers, revenue or measured fun',.72,6.53,11.9,.35,12,MUTED)
notes(s,'Availability source: OWNER_ACTIONS.md and submission/release-checklist.md. Local Practice and offline release do not establish public HTTPS or judge access. No live URL is invented. No new provider use, deployment, hosting purchase, repository visibility change, upload or event submission. Exhausted linked C/D/E QA allowance remains historical and is not a public-player allowance. RELEASE_NOT_LIVE_VERIFIED.')

output=ROOT/args.output
output.parent.mkdir(parents=True,exist_ok=True)
prs.save(output)
print(output)

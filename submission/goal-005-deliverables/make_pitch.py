"""Native Goal 005 presentation. Local files only; prior deliverables stay untouched."""
from pathlib import Path
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

ROOT = Path(__file__).resolve().parents[2]
STATUS = json.loads((Path(__file__).parent / 'status.json').read_text())
OUTPUT = ROOT / 'submission/Talk_Me_Home_Goal005_Pitch.pptx'
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
prs.core_properties.title = 'Talk Me Home — a Gallery worth exploring'
prs.core_properties.subject = STATUS['displayStatus']
prs.core_properties.author = 'Talk Me Home project'
PAPER, INK, GREEN, GOLD, MUTED = 'F4EFE2', '24382F', '557862', 'B58A46', '607066'

def box(s, x, y, w, h, fill):
    v = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    v.fill.solid(); v.fill.fore_color.rgb = RGBColor.from_string(fill); v.line.fill.background()
    return v

def text(s, copy, x, y, w, h, size=22, color=INK, bold=False, serif=False):
    v = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    f = v.text_frame; f.clear(); f.word_wrap = True
    f.margin_left = f.margin_right = f.margin_top = f.margin_bottom = 0
    for n, line in enumerate(copy.split('\n')):
        p = f.paragraphs[0] if n == 0 else f.add_paragraph()
        p.text = line; p.font.name = 'Liberation Serif' if serif else 'Liberation Sans'
        p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(9); p.line_spacing = 1.08
    return v

def picture(s, path, x, y, w, h):
    path = ROOT / path; iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih); dw, dh = iw * scale, ih * scale
    return s.shapes.add_picture(str(path), Inches(x + (w-dw)/2), Inches(y+(h-dh)/2), width=Inches(dw), height=Inches(dh))

def slide(n, kicker, title, sub=None):
    s = prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor.from_string(PAPER)
    box(s, .65, .5, .42, .05, GOLD)
    text(s, kicker.upper(), 1.2, .39, 11, .3, 11, MUTED, True)
    text(s, title, .65, .96, 12, .9, 35, serif=True)
    if sub: text(s, sub, .68, 1.88, 11.9, .55, 17, MUTED)
    text(s, STATUS['displayStatus'], .68, 7.03, 10.5, .25, 10, MUTED)
    text(s, f'{n:02d} / 06', 11.76, 7.02, .9, .25, 10, MUTED)
    return s

def notes(s, copy):
    s.notes_slide.notes_text_frame.text = copy

s = prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor.from_string(INK)
art = picture(s, 'submission/assets/cover-illustration-1600x900.png', 7.1, 0, 6.24, 7.5)
art.left, art.top, art.width, art.height = Inches(7.1), 0, Inches(6.24), Inches(7.5)
art.crop_left = .532
box(s, 0, 0, 7.15, 7.5, INK)
text(s, 'VOICE-LED CO-OP RESCUE', .76, .73, 6, .4, 13, 'CBD3BC', True)
text(s, 'Talk Me\nHome.', .73, 1.5, 6.1, 2.1, 59, 'F4EEDC', serif=True)
text(s, 'You have the map.\nPip has eyes and hands.', .78, 4.1, 5.9, 1.1, 24, 'EADBBC')
text(s, 'Explore through conversation.\nChoose the route. Confirm the action.', .78, 5.65, 5.9, .95, 18, 'DCE2CE')
text(s, STATUS['displayStatus'], .78, 7.04, 6, .25, 10, 'CBD3BC')
notes(s, 'Existing original cover illustration, not a gameplay screenshot. Goal 005 versioned local draft. Final results are controlled by status.json; no public link is invented.')

s = slide(2, 'The cooperation premise', 'Neither partner has the whole picture.', 'Three chapters: Cargo Bay → Relay Gallery → Return Dock.')
for x, title, body, foot, fill in [
    (.68, 'MISSION CONTROL', 'Read private documents.\nChoose a route.\nSet remote equipment.', 'Confirm each exact physical proposal.', 'E8E1CF'),
    (6.88, 'PIP', 'Observe nearby rooms.\nReport what is known.\nPropose a local action.', 'Looking and checking need no permission.', 'E4E9DD')]:
    box(s, x, 2.65, 5.77, 3.48, fill)
    text(s, title, x+.3, 2.98, 5.1, .4, 15, GREEN, True)
    text(s, body, x+.3, 3.63, 5.1, 1.75, 23)
    text(s, foot, x+.3, 5.62, 5.1, .47, 15, MUTED)
text(s, 'The server validates the current conditions. Spoken agreement does not press Confirm.', .8, 6.48, 11.9, .5, 17, GREEN)
notes(s, 'Sources: game/server/proposals.ts, sessions.ts and private records; existing shared contracts. Human documents and notes do not enter Pip context. Robot-local perception does not automatically position the human atlas. This is a normal gameplay split, not a developer-tools secrecy guarantee.')

s = slide(3, 'A Gallery worth exploring', 'Let a report change your plan.', 'Distinct emblems. Fixed compass. Circuit patterns. Your own route marks.')
picture(s, STATUS['uiCaptures']['wide'], .68, 2.53, 8.25, 4.22)
text(s, 'PLAN', 9.3, 2.75, 3.1, .35, 13, GREEN, True)
text(s, 'Mark a corridor.\nRevise the plan.', 9.3, 3.24, 3.14, .98, 20)
text(s, 'COMPARE', 9.3, 4.67, 3.1, .35, 13, GREEN, True)
text(s, 'Keep explored paths and\nreported obstructions distinct.', 9.3, 5.15, 3.14, 1.04, 18)
notes(s, 'Actual Goal 005 compiled Practice UI, artifacts/goal-005/ui/gallery-validation.json. Capture is deterministic offline gameplay, not real-model completion. Planned and explored routes are deliberate private player annotations. Both authored Gallery configurations and topology remain finite; no automatically solved safe path is drawn.')

s = slide(4, 'Useful field reports', 'Exact words, with their source and age.', 'A report is a communicated claim, not a live sensor reading.')
picture(s, STATUS['uiCaptures']['report'], .68, 2.6, 7.65, 4.22)
text(s, 'Keep the quote.', 8.76, 2.86, 3.82, .48, 24, serif=True)
text(s, 'Attach it to a room or corridor.\nThe association belongs to you.', 8.76, 3.43, 3.82, .9, 17, MUTED)
text(s, 'Recheck what changed.', 8.76, 4.74, 3.82, .48, 24, serif=True)
text(s, 'Relay changes age gate notes.\nStable clues remain historical quotes.\nAsk for a fresh surroundings report.', 8.76, 5.3, 3.82, 1.23, 16, MUTED)
notes(s, 'Actual Practice screenshot at CSS 200% zoom, not a visual mockup. Exact saved messages and original report time/source are preserved. Report freshness is server-stamped at receipt; a Relay change and change-back cannot revive a gate reading. Selected quick requests are explicitly labelled Selected request, never microphone speech. Private associations do not enter robot context.')

s = slide(5, 'A companion that helps you orient', '“I moved” should not be the end of the exchange.', 'Goal 005 connects a confirmed arrival to a useful local report.')
for x, title, body, fill in [
    (.68, '1  CONFIRMED ARRIVAL', 'One confirmed action.\nA scoped local observation.\nIts source and time retained.', 'E4E9DD'),
    (4.9, '2  USEFUL ORIENTATION', 'Report the nearby emblem.\nName visible directions.\nKeep the response bounded.', 'E8E1CF'),
    (9.12, '3  A HUMAN CHOICE', 'Compare with the atlas.\nCheck an uncertain branch.\nReplan without resetting.', 'E4E9DD')]:
    box(s, x, 2.79, 3.53, 3.17, fill)
    text(s, title, x+.23, 3.11, 3.1, .38, 12, GREEN, True)
    text(s, body, x+.23, 3.83, 3.07, 1.8, 18)
text(s, 'Pause, Interrupt and End still win. A missing report leaves room to ask again.', .8, 6.4, 11.9, .53, 19, GREEN)
notes(s, 'Implemented Goal 005 product design; schema/adapter and compiled injected-provider tests establish local behavior, not a guarantee of real-model compliance. Arrival data is robot-eligible only and has no hidden global graph. Existing one-shot scheduler preserves safe tool/playback boundaries and cancellation. Final live efficacy must be supplied from completed batch evidence.')

s = slide(6, 'Evidence and availability', STATUS['acceptanceHeadline'])
box(s, .68, 2.25, 7.1, 3.83, 'E4E9DD')
text(s, 'CURRENT EVIDENCE', .98, 2.62, 6.35, .35, 13, GREEN, True)
text(s, STATUS['acceptanceCopy'], .98, 3.21, 6.36, 2.22, 22)
text(s, 'LOCAL PROTOTYPE', 8.2, 2.62, 4.43, .35, 13, GREEN, True)
text(s, 'Playable Rescue and Training.\nNo public demo is deployed.\nNo upload or event entry is submitted.', 8.2, 3.23, 4.38, 1.9, 20)
text(s, 'Synthetic QA is not a human\nplaytest or physical-device test.', 8.2, 5.47, 4.38, .88, 16, MUTED)
text(s, f"Evidence: {STATUS['evidencePath']}", .8, 6.53, 11.9, .3, 12, MUTED)
notes(s, 'This slide must be updated from actual final-candidate results. Do not convert Practice, fake-peer tests, selected requests or historical failed calls into a Voice completion. No claim about human enjoyment, physical audio quality, revenue, customers or statistical reliability. Public URLs intentionally remain unset. ' + json.dumps(STATUS))

prs.save(OUTPUT)
print(OUTPUT)
